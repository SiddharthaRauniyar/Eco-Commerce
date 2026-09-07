"""Staff reporting checks for access, figures, and portable downloads."""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import CustomUser
from apps.categories.models import Category
from apps.orders.models import Order, OrderItem
from apps.products.models import Product


class ReportsTests(TestCase):
    def setUp(self):
        category = Category.objects.create(name="Reports", slug="reports")
        self.product = Product.objects.create(
            category=category,
            name="Report product",
            slug="report-product",
            sku="REPORT-001",
            description="A product used for analytics tests.",
            base_price=Decimal("25.00"),
            stock_quantity=1,
            reorder_level=2,
            status=Product.Status.ACTIVE,
        )
        self.customer = CustomUser.objects.create_user("analytics-customer@example.com", "safe-test-password")
        self.staff = CustomUser.objects.create_superuser("analytics-staff@example.com", "safe-test-password")
        self.order = Order.objects.create(
            order_number="SC-REPORT-0001",
            user=self.customer,
            billing_address_snapshot={},
            shipping_address_snapshot={},
            status=Order.Status.CONFIRMED,
            payment_status=Order.PaymentStatus.PAID,
            currency="USD",
            subtotal=Decimal("50.00"),
            shipping_total=Decimal("5.00"),
            grand_total=Decimal("55.00"),
        )
        OrderItem.objects.create(
            order=self.order,
            product=self.product,
            product_name=self.product.name,
            sku=self.product.sku,
            unit_price=Decimal("25.00"),
            quantity=2,
            line_total=Decimal("50.00"),
        )
        Order.objects.filter(pk=self.order.pk).update(placed_at=timezone.now() - timedelta(days=1))

    def test_staff_reports_show_paid_sales_and_no_customer_pii(self):
        self.client.force_login(self.staff)

        response = self.client.get(reverse("analytics:reports"), {"range": 7})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "USD 55.00")
        self.assertContains(response, self.product.name)
        self.assertNotContains(response, self.customer.email)

    def test_export_formats_are_parseable_and_exclude_customer_email(self):
        self.client.force_login(self.staff)

        csv_response = self.client.get(reverse("analytics:export", args=("csv",)), {"range": 7})
        xlsx_response = self.client.get(reverse("analytics:export", args=("xlsx",)), {"range": 7})
        pdf_response = self.client.get(reverse("analytics:export", args=("pdf",)), {"range": 7})

        self.assertEqual(csv_response.status_code, 200)
        self.assertIn("text/csv", csv_response["Content-Type"])
        self.assertIn(b"Report product", csv_response.content)
        self.assertNotIn(self.customer.email.encode(), csv_response.content)
        self.assertEqual(load_workbook(BytesIO(xlsx_response.content)).sheetnames[0], "Summary")
        self.assertTrue(pdf_response.content.startswith(b"%PDF"))

    def test_non_staff_cannot_access_reports_or_downloads(self):
        self.client.force_login(self.customer)

        self.assertEqual(self.client.get(reverse("analytics:reports")).status_code, 403)
        self.assertEqual(self.client.get(reverse("analytics:export", args=("csv",))).status_code, 404)
