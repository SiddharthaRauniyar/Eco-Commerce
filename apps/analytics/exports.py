"""Download renderers for compact, non-PII staff reports."""

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas


def _xlsx_value(value):
    """Excel accepts naive dates, not Django's timezone-aware timestamps."""

    return value.replace(tzinfo=None) if getattr(value, "tzinfo", None) else value


def report_tables(data: dict) -> list[tuple[str, list[str], list[list]]]:
    """Present each permitted report area as a portable table."""

    tables = [("Summary", ["Metric", "Value"], [[item["label"], item["value"]] for item in data["summary"]])]
    if data["can_orders"]:
        tables.extend(
            [
                ("Daily sales", ["Date", "Orders", "Revenue (USD)"], [[row["period"], row["orders"], row["revenue"]] for row in data["daily"]]),
                ("Weekly sales", ["Week", "Orders", "Revenue (USD)"], [[row["label"], row["orders"], row["revenue"]] for row in data["weekly"]]),
                ("Monthly sales", ["Month", "Orders", "Revenue (USD)"], [[row["label"], row["orders"], row["revenue"]] for row in data["monthly"]]),
                ("Product performance", ["Product", "SKU", "Units sold", "Revenue (USD)"], [[row["product_name"], row["sku"], row["units_sold"], row["revenue"]] for row in data["products"]]),
            ]
        )
    if data["can_products"]:
        tables.append(("Inventory", ["Item", "SKU", "On hand", "Reorder level"], [[row["name"], row["sku"], row["stock"], row["reorder_level"]] for row in data["inventory"]]))
    if data["can_coupons"]:
        tables.append(("Coupon performance", ["Coupon", "Orders", "Discount given (USD)"], [[row["coupon__code"], row["orders"], row["discount_total"]] for row in data["coupons"]]))
    if data["can_security"]:
        tables.append(("Security events", ["Severity", "Events"], [[row["severity"], row["events"]] for row in data["security"]]))
    return tables


def csv_bytes(data: dict) -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["Secure Commerce report", f"{data['start']} to {data['end']}"])
    for title, headers, rows in report_tables(data):
        writer.writerow([])
        writer.writerow([title])
        writer.writerow(headers)
        writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def xlsx_bytes(data: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="4F49C8")
    header_font = Font(color="FFFFFF", bold=True)
    for title, headers, rows in report_tables(data):
        sheet = workbook.create_sheet(title[:31])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill, cell.font = header_fill, header_font
        for row in rows:
            sheet.append([_xlsx_value(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_index, header in enumerate(headers, 1):
            widest_value = max(
                [len(str(row[column_index - 1])) for row in rows if len(row) >= column_index],
                default=0,
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(len(header), widest_value) + 3, 36
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if "Revenue" in headers[cell.column - 1] or "Discount" in headers[cell.column - 1]:
                    cell.number_format = '"USD" #,##0.00'
                elif "Date" in headers[cell.column - 1] or "Week" in headers[cell.column - 1] or "Month" in headers[cell.column - 1]:
                    cell.number_format = "yyyy-mm-dd"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def pdf_bytes(data: dict) -> bytes:
    """Render an at-a-glance print report, keeping detailed tables in exports."""

    stream = BytesIO()
    pdf = canvas.Canvas(stream, pagesize=A4, pageCompression=0)
    width, height = A4
    left, right, y = 20 * mm, width - 20 * mm, height - 20 * mm
    ink, muted, brand, rule = HexColor("#20233a"), HexColor("#687087"), HexColor("#4f49c8"), HexColor("#dfe3f4")

    def page_header() -> float:
        pdf.setFillColor(brand)
        pdf.rect(0, height - 30 * mm, width, 30 * mm, fill=1, stroke=0)
        pdf.setFillColor(HexColor("#ffffff"))
        pdf.setFont("Helvetica-Bold", 18)
        pdf.drawString(left, height - 16 * mm, "SECURE COMMERCE")
        pdf.setFont("Helvetica", 9)
        pdf.drawString(left, height - 22 * mm, f"Operations report - {data['start']} to {data['end']}")
        return height - 44 * mm

    def section(title: str, headers: list[str], rows: list[list], current_y: float) -> float:
        if current_y < 55 * mm:
            pdf.showPage()
            current_y = page_header()
        pdf.setFillColor(ink)
        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawString(left, current_y, title)
        current_y -= 6 * mm
        pdf.setFillColor(HexColor("#eef0ff"))
        pdf.rect(left, current_y - 5 * mm, right - left, 6 * mm, fill=1, stroke=0)
        pdf.setFillColor(ink)
        pdf.setFont("Helvetica-Bold", 8)
        width_per_column = (right - left) / len(headers)
        for index, header in enumerate(headers):
            pdf.drawString(left + 2 * mm + index * width_per_column, current_y - 3.8 * mm, str(header)[:22])
        current_y -= 8 * mm
        pdf.setFont("Helvetica", 8)
        for row in rows[:10]:
            if current_y < 30 * mm:
                pdf.showPage()
                current_y = page_header()
            pdf.setStrokeColor(rule)
            pdf.line(left, current_y - 2 * mm, right, current_y - 2 * mm)
            for index, value in enumerate(row):
                text = str(value)
                max_chars = 24 if index == 0 else 16
                pdf.drawString(left + 2 * mm + index * width_per_column, current_y - 6 * mm, text[:max_chars])
            current_y -= 7 * mm
        return current_y - 6 * mm

    y = page_header()
    for title, headers, rows in report_tables(data):
        y = section(title, headers, rows, y)
    pdf.setStrokeColor(rule)
    pdf.line(left, 16 * mm, right, 16 * mm)
    pdf.setFillColor(muted)
    pdf.setFont("Helvetica", 8)
    pdf.drawString(left, 10 * mm, "Secure Commerce staff report")
    pdf.save()
    return stream.getvalue()
